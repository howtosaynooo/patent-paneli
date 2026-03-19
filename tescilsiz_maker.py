import os
import time
from urllib.parse import quote

from openpyxl import Workbook, load_workbook


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_EXCEL = os.path.join(BASE_DIR, "output.xlsx")
TESCILSIZ_EXCEL = os.path.join(BASE_DIR, "Tescilsiz.xlsx")
LOGS_DIR = os.path.join(BASE_DIR, "logs")

STOP_REQUESTED = False
PROGRESS = {"total": 0, "processed": 0}


def log(message: str) -> None:
    ts = time.strftime("%H:%M:%S")
    text = f"[{ts}] {message}"
    print(text)

    os.makedirs(LOGS_DIR, exist_ok=True)
    log_path = os.path.join(LOGS_DIR, "tescilsiz_maker.log")
    try:
        with open(log_path, "a", encoding="utf-8") as f:
            f.write(text + "\n")
    except OSError:
        pass


def trendyol_link(brand_name: str) -> str:
    encoded = quote(brand_name, safe="")
    return f"https://www.trendyol.com/sr?q={encoded}&qt={encoded}&st={encoded}&os=1"


def load_output() -> list[dict]:
    if not os.path.exists(OUTPUT_EXCEL):
        log(f"output.xlsx bulunamadı: {OUTPUT_EXCEL}")
        return []

    wb = load_workbook(OUTPUT_EXCEL)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        marka = row[0]
        patent = row[1]
        kategori = row[2]
        if not marka:
            continue
        rows.append({
            "marka": str(marka).strip(),
            "patent": patent,
            "kategori_id": str(kategori).strip() if kategori else "",
        })
    log(f"output.xlsx'ten {len(rows)} kayıt okundu.")
    return rows


def load_existing_brands() -> set:
    brands = set()
    if not os.path.exists(TESCILSIZ_EXCEL):
        return brands
    try:
        wb = load_workbook(TESCILSIZ_EXCEL)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0]:
                brands.add(str(row[0]).strip())
    except Exception:
        pass
    return brands


def main():
    global STOP_REQUESTED
    STOP_REQUESTED = False

    rows = load_output()
    if not rows:
        log("İşlenecek kayıt bulunamadı.")
        return

    tescilsiz = [r for r in rows if r["patent"] is False or r["patent"] == False]
    log(f"Tescilsiz marka sayısı: {len(tescilsiz)} / {len(rows)}")

    existing = load_existing_brands()
    pending = [r for r in tescilsiz if r["marka"] not in existing]
    log(f"Yeni eklenecek: {len(pending)} ({len(tescilsiz) - len(pending)} zaten mevcut)")

    PROGRESS["total"] = len(pending)
    PROGRESS["processed"] = 0

    if not pending:
        log("Eklenecek yeni tescilsiz marka yok.")
        return

    file_exists = os.path.exists(TESCILSIZ_EXCEL)
    if file_exists:
        wb = load_workbook(TESCILSIZ_EXCEL)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["Marka", "Kategori ID", "Trendyol Link"])

    for r in pending:
        if STOP_REQUESTED:
            log("Durdurma isteği alındı.")
            break
        ws.append([r["marka"], r["kategori_id"], trendyol_link(r["marka"])])
        PROGRESS["processed"] += 1

    wb.save(TESCILSIZ_EXCEL)
    log(f"Tescilsiz.xlsx güncellendi. Toplam eklenen: {PROGRESS['processed']}")


if __name__ == "__main__":
    main()
