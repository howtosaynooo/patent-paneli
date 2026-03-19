"""
eleyici.py
Tescilsiz.xlsx'teki her marka için, kaydedildiği kategori+gender kombinasyonunda
Trendyol top-ranking API'sini sorgular; markanın ürünlerini filtreler ve
değerlendirme puanı + sayısını elenmis.xlsx'e yazar.
"""
import os
import time
import random

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from openpyxl import Workbook, load_workbook


BASE_URL = "https://apigw.trendyol.com"
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
TESCILSIZ_EXCEL = os.path.join(BASE_DIR, "Tescilsiz.xlsx")
ELENMIS_EXCEL = os.path.join(BASE_DIR, "elenmis.xlsx")
LOGS_DIR = os.path.join(BASE_DIR, "logs")

STOP_REQUESTED = False
PROGRESS = {"total": 0, "processed": 0}

HEADERS = {
    "User-Agent": "Mozilla/5.0 (X11; Linux x86_64; rv:128.0) Gecko/20100101 Firefox/128.0",
    "Accept": "application/json, text/plain, */*",
}

session = requests.Session()
retries = Retry(
    total=3,
    backoff_factor=0.5,
    status_forcelist=[429, 500, 502, 503, 504],
    allowed_methods=["GET"],
)
adapter = HTTPAdapter(max_retries=retries)
session.mount("https://", adapter)
session.mount("http://", adapter)


def log(message: str) -> None:
    ts = time.strftime("%H:%M:%S")
    text = f"[{ts}] {message}"
    print(text)
    os.makedirs(LOGS_DIR, exist_ok=True)
    try:
        with open(os.path.join(LOGS_DIR, "eleyici.log"), "a", encoding="utf-8") as f:
            f.write(text + "\n")
    except OSError:
        pass


def random_sleep(min_s=0.3, max_s=0.8):
    time.sleep(random.uniform(min_s, max_s))


def load_tescilsiz() -> list[dict]:
    if not os.path.exists(TESCILSIZ_EXCEL):
        log(f"Tescilsiz.xlsx bulunamadı: {TESCILSIZ_EXCEL}")
        return []
    wb = load_workbook(TESCILSIZ_EXCEL)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        rows.append({
            "marka": str(row[0]).strip(),
            "kategori_id": str(row[1]).strip() if row[1] else "",
            "link": str(row[2]).strip() if row[2] else "",
        })
    log(f"Tescilsiz.xlsx'ten {len(rows)} marka yüklendi.")
    return rows


def load_processed_brands() -> set:
    brands = set()
    if not os.path.exists(ELENMIS_EXCEL):
        return brands
    try:
        wb = load_workbook(ELENMIS_EXCEL)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if row[0]:
                brands.add(str(row[0]).strip())
    except Exception:
        pass
    return brands


def append_results(results: list[dict]) -> None:
    if not results:
        return
    file_exists = os.path.exists(ELENMIS_EXCEL)
    if file_exists:
        wb = load_workbook(ELENMIS_EXCEL)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["Marka", "Değerlendirme Sayısı", "Değerlendirme Puanı", "Link"])
    for r in results:
        ws.append([r["marka"], r["rev_count"], r["rating"], r["link"]])
    wb.save(ELENMIS_EXCEL)
    log(f"elenmis.xlsx güncellendi ({len(results)} yeni satır).")


def parse_kategori_id(raw: str) -> tuple[int | None, str]:
    """'28-1' → (28, '1')   |   '28' → (28, '1')"""
    raw = raw.strip()
    if "-" in raw:
        parts = raw.split("-", 1)
        cat = int(parts[0]) if parts[0].isdigit() else None
        gender = parts[1] if parts[1].isdigit() else "1"
        return cat, gender
    return (int(raw) if raw.isdigit() else None), "1"


def fetch_top_ranking(category_id: int, gender: str) -> list[dict]:
    """apigw.trendyol.com top ranking API'sini çağır, contents listesini döndür."""
    url = f"{BASE_URL}/discovery-web-websfxcategorytopranking-santral/topRankingContents"
    params = {
        "categoryId": category_id,
        "rankingType": "bestSeller",
        "gender": gender,
        "page": "1",
        "channelId": "1",
    }
    try:
        resp = session.get(url, headers=HEADERS, params=params, timeout=15)
        resp.raise_for_status()
        return resp.json().get("result", {}).get("contents", [])
    except Exception as e:
        log(f"API hatası (kategori={category_id}, gender={gender}): {e}")
        return []


def extract_brand_products(contents: list[dict], brand_name: str) -> list[dict]:
    """contents listesinden marka adına göre filtreleyip rating bilgisini döndür."""
    results = []
    for item in contents:
        item_brand = (item.get("brand") or {}).get("name", "")
        if item_brand.strip().lower() != brand_name.strip().lower():
            continue

        rating_data = item.get("ratingScore") or {}
        rating = rating_data.get("averageRating", "")
        rev_count = rating_data.get("totalCount", "")

        # Ürün linki için id'den basit arama URL'si kullan
        product_id = item.get("id", "")
        link = f"https://www.trendyol.com/p-{product_id}" if product_id else ""

        results.append({
            "rating": str(round(rating, 1)) if isinstance(rating, float) else str(rating),
            "rev_count": str(rev_count),
            "link": link,
        })

        if len(results) >= 4:
            break

    return results


def main():
    global STOP_REQUESTED
    STOP_REQUESTED = False

    rows = load_tescilsiz()
    if not rows:
        log("İşlenecek marka bulunamadı.")
        return

    processed_brands = load_processed_brands()
    pending = [r for r in rows if r["marka"] not in processed_brands]

    PROGRESS["total"] = len(rows)
    PROGRESS["processed"] = len(rows) - len(pending)

    log(f"Toplam {len(rows)} marka, {len(pending)} işlenecek ({len(rows) - len(pending)} zaten işlenmiş).")

    if not pending:
        log("Tüm markalar zaten işlenmiş.")
        return

    batch: list[dict] = []

    for entry in pending:
        if STOP_REQUESTED:
            log("Durdurma isteği alındı.")
            break

        brand = entry["marka"]
        cat_id, gender = parse_kategori_id(entry["kategori_id"])

        if cat_id is None:
            log(f"{brand}: Geçersiz kategori ID '{entry['kategori_id']}', atlanıyor.")
            PROGRESS["processed"] += 1
            continue

        contents = fetch_top_ranking(cat_id, gender)
        products = extract_brand_products(contents, brand)

        if products:
            for p in products:
                batch.append({
                    "marka": brand,
                    "rev_count": p["rev_count"],
                    "rating": p["rating"],
                    "link": entry["link"],
                })
            log(f"{brand}: {len(products)} ürün bulundu (kategori={cat_id}, gender={gender}).")
        else:
            # API'de bulunamadı — boş satır yaz ki tekrar işlenmesin
            batch.append({
                "marka": brand,
                "rev_count": "",
                "rating": "",
                "link": entry["link"],
            })
            log(f"{brand}: Bu kategoride ürün bulunamadı.")

        PROGRESS["processed"] += 1
        random_sleep()

    if batch:
        append_results(batch)

    log(f"Eleyici tamamlandı. İşlenen: {PROGRESS['processed']} / {PROGRESS['total']}")


if __name__ == "__main__":
    main()
