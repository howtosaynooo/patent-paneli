import os
import time
import random
from threading import Lock
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from openpyxl import Workbook, load_workbook

from selenium import webdriver
from selenium.webdriver.firefox.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.common.exceptions import NoSuchElementException, TimeoutException, ElementClickInterceptedException
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC


BASE_URL = "https://apigw.trendyol.com"
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CACHE_DIR = os.path.join(BASE_DIR, "cache")
CATEGORIES_FILE = os.path.join(CACHE_DIR, "categories_with_products.txt")
OUTPUT_EXCEL = os.path.join(BASE_DIR, "output.xlsx")
TESCILSIZ_EXCEL = os.path.join(BASE_DIR, "Tescilsiz.xlsx")
LOGS_DIR = os.path.join(BASE_DIR, "logs")

NUM_WORKERS = 10
UPDATE_INTERVAL = 10  # Kaç kayıt biriktirince Excel'e yazılacak


visited_brands = set()
visited_lock = Lock()

# Durdurma isteği; panel "Durdur" tıklanınca True yapar
STOP_REQUESTED = False

# Basit ilerleme bilgisi; panel.py bu sözlüğü okuyabilir
PROGRESS = {
    "total_pairs": 0,
    "processed_pairs": 0,
    "total_brands": 0,
    "processed_brands": 0,
}


def log(message: str) -> None:
    ts = time.strftime("%H:%M:%S")
    text = f"[{ts}] {message}"
    print(text)

    os.makedirs(LOGS_DIR, exist_ok=True)
    log_path = os.path.join(LOGS_DIR, "patent_worker.log")
    try:
        with open(log_path, "a", encoding="utf-8") as f:
            f.write(text + "\n")
    except OSError:
        pass


# HTTP session + retry
session = requests.Session()
retries = Retry(
    total=5,
    backoff_factor=0.5,
    status_forcelist=[429, 500, 502, 503, 504],
    allowed_methods=["GET", "POST"],
)
adapter = HTTPAdapter(max_retries=retries)
session.mount("https://", adapter)
session.mount("http://", adapter)


def load_categories(path: str = CATEGORIES_FILE) -> list[tuple[int, str]]:
    if not os.path.exists(path):
        log(f"Kategori cache dosyası bulunamadı: {path}")
        return []

    pairs: list[tuple[int, str]] = []
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            s = line.strip()
            if not s:
                continue
            if "," in s:
                cid_str, gender_str = [p.strip() for p in s.split(",", 1)]
                if cid_str.isdigit() and gender_str.isdigit():
                    pairs.append((int(cid_str), gender_str))
            elif s.isdigit():
                pairs.append((int(s), "1"))

    seen = set()
    out = []
    for cid, gender in pairs:
        key = (cid, gender)
        if key not in seen:
            seen.add(key)
            out.append(key)

    PROGRESS["total_pairs"] = len(out)
    log(f"Yüklendi: {len(out)} kategori-gender çifti (cache)")
    return out


def trendyol_link(brand_name: str) -> str:
    from urllib.parse import quote
    encoded = quote(brand_name, safe="")
    return f"https://www.trendyol.com/sr?q={encoded}&qt={encoded}&st={encoded}&os=1"


def _load_existing_brands(ws) -> set:
    """Çalışma sayfasındaki mevcut marka adlarını (A kolonu) döndürür."""
    brands = set()
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0]:
            brands.add(str(row[0]))
    return brands


def excel_export(names, answers, category_ids):
    file_path = OUTPUT_EXCEL
    file_exists = os.path.exists(file_path)

    if file_exists:
        wb = load_workbook(file_path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["Marka", "Patent Var mı?", "Kategori ID"])

    existing = _load_existing_brands(ws)
    added = 0
    for n, a, c in zip(names, answers, category_ids):
        if n not in existing:
            ws.append([n, a, c])
            existing.add(n)
            added += 1

    wb.save(file_path)
    log(f"output.xlsx güncellendi ({added} yeni kayıt).")


def excel_export_tescilsiz(names, answers, category_ids):
    """Patent bulunamayan markaları Tescilsiz.xlsx'e yazar (mükerrer önlenmiş)."""
    file_path = TESCILSIZ_EXCEL
    file_exists = os.path.exists(file_path)

    if file_exists:
        wb = load_workbook(file_path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["Marka", "Kategori ID", "Trendyol Link"])

    existing = _load_existing_brands(ws)
    added = 0
    for n, a, c in zip(names, answers, category_ids):
        if not a and n not in existing:
            ws.append([n, c, trendyol_link(n)])
            existing.add(n)
            added += 1

    wb.save(file_path)
    log(f"Tescilsiz.xlsx güncellendi ({added} yeni kayıt).")


def random_sleep(min_seconds=0.3, max_seconds=0.8):
    time.sleep(random.uniform(min_seconds, max_seconds))


def close_popup(driver):
    try:
        popup_button = WebDriverWait(driver, 5).until(
            EC.element_to_be_clickable((By.XPATH, '//*[@id="__next"]/div/section/div/span'))
        )

        driver.execute_script("arguments[0].scrollIntoView(true);", popup_button)
        random_sleep()

        popup_button.click()
        log("Pop-up kapatıldı.")
        random_sleep()

    except TimeoutException:
        log("Pop-up bulunamadı, devam ediliyor...")
    except Exception as e:
        log(f"Pop-up kapatılamadı: {str(e)}")


def get_products_from_top_ranking(params):
    headers = {
        "User-Agent": "Mozilla/5.0 (X11; Linux x86_64; rv:128.0) Gecko/20100101 Firefox/128.0",
        "Accept": "application/json, text/plain, */*",
    }

    url = "/discovery-web-websfxcategorytopranking-santral/topRankingContents"
    full_url = f"{BASE_URL}{url}"

    response = session.get(full_url, headers=headers, params=params, timeout=15)
    response.raise_for_status()
    return response.json().get("result", {})


def ask_for_patent(driver, brand_name):
    isExist = False
    try:
        driver.get("https://www.turkpatent.gov.tr/arastirma-yap")
        close_popup(driver)
        random_sleep()

        WebDriverWait(driver, 10).until(
            EC.presence_of_element_located((By.XPATH, '//input[@placeholder="Marka Adı"]'))
        )

        input_field = driver.find_element(By.XPATH, '//input[@placeholder="Marka Adı"]')

        driver.execute_script("arguments[0].scrollIntoView(true);", input_field)
        random_sleep()
        input_field.send_keys(brand_name)
        input_field.send_keys(Keys.RETURN)
        random_sleep()

        WebDriverWait(driver, 10).until(
            EC.presence_of_element_located((By.XPATH, '//input[@class="jss45" and @value="isEqual"]'))
        )

        radio_button = driver.find_element(By.XPATH, '//input[@class="jss45" and @value="isEqual"]')

        # Radio buton tıklamasında ara sıra overlay çakışması olabiliyor; birkaç kez yeniden dene
        clicked = False
        for attempt in range(3):
            try:
                driver.execute_script("arguments[0].scrollIntoView(true);", radio_button)
                random_sleep()
                radio_button.click()
                random_sleep()
                clicked = True
                break
            except ElementClickInterceptedException:
                print(f"{brand_name}: Radio buton tıklanamadı, tekrar denenecek ({attempt + 1}/3).")
                random_sleep()

        if not clicked:
            print(f"{brand_name}: Radio buton 3 denemede de tıklanamadı, marka atlanıyor.")
            return isExist

        WebDriverWait(driver, 10).until(
            EC.presence_of_element_located(
                (By.CSS_SELECTOR, ".MuiButtonBase-root.MuiButton-root.MuiButton-contained.MuiButton-containedSecondary")
            )
        )

        search_button = driver.find_element(
            By.CSS_SELECTOR, ".MuiButtonBase-root.MuiButton-root.MuiButton-contained.MuiButton-containedSecondary"
        )
        search_button.click()
        random_sleep()

        try:
            WebDriverWait(driver, 10).until(
                EC.presence_of_element_located((By.CSS_SELECTOR, "#results > thead > tr > th:nth-child(2) > span"))
            )
            isExist = True
            log(f"{brand_name}: Patent bulundu.")
        except TimeoutException:
            log(f"{brand_name}: Patent bulunamadı.")

    except (NoSuchElementException, TimeoutException, ElementClickInterceptedException) as e:
        log(f"{brand_name}: Hata: {str(e)}")

    return isExist


def create_driver():
    options = Options()
    options.add_argument("--headless")
    return webdriver.Firefox(options=options)


def process_brands(driver, brands, results, category_id):
    for brand in brands:
        if STOP_REQUESTED:
            return
        with visited_lock:
            if brand in visited_brands:
                continue
            visited_brands.add(brand)

        isExist = ask_for_patent(driver, brand)
        results.append((brand, isExist, category_id))
        PROGRESS["processed_brands"] += 1


def distribute_brands(brands, max_workers):
    n = len(brands)
    if n == 0:
        return []

    worker_count = min(max_workers, n)
    chunk_size = (n + worker_count - 1) // worker_count

    distributed = [
        brands[i * chunk_size : (i + 1) * chunk_size]
        for i in range(worker_count)
    ]
    return distributed


def main():
    global STOP_REQUESTED
    STOP_REQUESTED = False

    category_pairs = load_categories()
    if not category_pairs:
        log("İşlenecek kategori-gender çifti bulunamadı. Önce category_finder.py çalıştırmalısın.")
        return

    PROGRESS["processed_pairs"] = 0
    PROGRESS["processed_brands"] = 0

    names = []
    answers = []
    cat_ids = []
    update_counter = 0

    drivers = [create_driver() for _ in range(NUM_WORKERS)]

    try:
        for idx, (categoryId, gender) in enumerate(category_pairs, start=1):
            if STOP_REQUESTED:
                log("Durdurma isteği alındı, patent sorgulama durduruluyor.")
                break

            PROGRESS["processed_pairs"] = idx
            log(f"Kategori {categoryId} (gender={gender}) için sorgulama başladı.")
            params = {
                "categoryId": categoryId,
                "rankingType": "bestSeller",
                "gender": gender,
                "page": "1",
                "channelId": "1",
            }

            try:
                res = get_products_from_top_ranking(params)
            except Exception as e:
                log(f"Kategori {categoryId} (gender={gender}) için API hatası: {e}")
                continue

            if res and "contents" in res and len(res["contents"]) > 0:
                contents = res["contents"]
                brands = list({con["brand"]["name"] for con in contents})

                if not brands:
                    log(f"Kategori {categoryId} (gender={gender}) için marka bulunamadı.")
                    continue

                distributed = distribute_brands(brands, NUM_WORKERS)
                worker_count = len(distributed)
                results = [[] for _ in range(worker_count)]

                with ThreadPoolExecutor(max_workers=worker_count) as executor:
                    futures = [
                        executor.submit(
                            process_brands,
                            drivers[i],
                            distributed[i],
                            results[i],
                            categoryId,
                        )
                        for i in range(worker_count)
                    ]

                    for future in as_completed(futures):
                        future.result()

                for result in results:
                    for brand, isExist, cid in result:
                        names.append(brand)
                        answers.append(isExist)
                        # Kategori ID'yi gender ile birlikte yaz (ör. "55-1")
                        cat_ids.append(f"{cid}-{gender}")
                        update_counter += 1

                        if update_counter % UPDATE_INTERVAL == 0:
                            excel_export(names, answers, cat_ids)
                            excel_export_tescilsiz(names, answers, cat_ids)
                            names, answers, cat_ids = [], [], []

                log(f"Kategori {categoryId} (gender={gender}) için sorgulama tamamlandı.")
            else:
                log(f"Kategori {categoryId} (gender={gender}) için ürün bulunamadı veya API yanıtı geçersiz.")

    finally:
        for d in drivers:
            try:
                d.quit()
            except Exception:
                pass

        if names or answers:
            excel_export(names, answers, cat_ids)
            excel_export_tescilsiz(names, answers, cat_ids)

        log("Patent sorgulama işlemi tamamlandı.")


if __name__ == "__main__":
    main()